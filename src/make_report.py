"""Build the assignment report (.docx) from real pipeline results.

Reads MLflow runs, EDA outputs and dataset statistics, embeds plots and any
screenshots found in screenshots/, and marks missing screenshots clearly.
Run:  python -m src.make_report
"""
import glob
import json
import os
import re
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import mlflow  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

from src.data import FEATURES, NUMERIC, load_clean, load_raw  # noqa: E402

# ---- edit these two lines if needed ---------------------------------------
STUDENT_NAME = "Siva Sankar K"
STUDENT_ID = "s2025ae05092"
# ---------------------------------------------------------------------------
REPO_URL = "https://github.com/shankarece/mlops-heart-disease"
ROOT = os.path.join(os.path.dirname(__file__), "..")
EDA = os.path.join(ROOT, "reports", "eda")
SHOTS = os.path.join(ROOT, "screenshots")
OUT = os.path.join(ROOT, "reports", "MLOps_Assignment1_Report.docx")
missing = []


def p(doc, text, bold=False, italic=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold, run.italic = bold, italic
    return para


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.add_paragraph()


def picture(doc, path, caption, width=6.0):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True


def shot(doc, filename, caption):
    path = os.path.join(SHOTS, filename)
    if os.path.exists(path):
        picture(doc, path, caption)
    else:
        missing.append(filename)
        para = doc.add_paragraph()
        run = para.add_run(f"[INSERT SCREENSHOT: {filename} - {caption}]")
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def draw_architecture(path):
    fig, ax = plt.subplots(figsize=(11, 4.2))
    ax.axis("off")
    w, h = 0.12, 0.22
    top, bot = 0.62, 0.16
    cols = [0.02, 0.19, 0.36, 0.53, 0.70, 0.86]
    boxes = {
        "uci": (cols[0], top, "UCI dataset"),
        "clean": (cols[1], top, "Clean +\npreprocess"),
        "train": (cols[2], top, "Train LR / RF\n(CV + grid)"),
        "model": (cols[3], top, "Model\npipeline"),
        "api": (cols[4], top, "FastAPI\n/predict"),
        "mon": (cols[5], top, "Prometheus\n+ Grafana"),
        "mlf": (cols[2], bot, "MLflow\ntracking"),
        "ci": (cols[3], bot, "GitHub Actions\nCI"),
        "dock": (cols[4], bot, "Docker\nimage"),
        "k8s": (cols[5], bot, "Kubernetes\n(LoadBalancer)"),
    }
    for x, y, label in boxes.values():
        ax.add_patch(plt.Rectangle((x, y), w, h, fc="#dbe9f6", ec="#2b5c8a", lw=1.5))
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=9)

    def arrow(a, b, style="-"):
        xa, ya, _ = boxes[a]
        xb, yb, _ = boxes[b]
        if ya == yb:  # horizontal: right edge of a -> left edge of b
            p1, p2 = (xa + w, ya + h / 2), (xb, yb + h / 2)
        else:         # vertical: bottom of a -> top of b
            p1, p2 = (xa + w / 2, ya), (xb + w / 2, yb + h)
        ax.annotate("", xy=p2, xytext=p1, arrowprops=dict(
            arrowstyle="->", lw=1.4, color="#333", linestyle=style))

    for a, b in [("uci", "clean"), ("clean", "train"), ("train", "model"),
                 ("model", "api"), ("api", "mon"), ("train", "mlf"),
                 ("api", "dock"), ("dock", "k8s")]:
        arrow(a, b)
    arrow("ci", "dock", "--")
    ax.text(cols[3] + w + 0.025, bot + h / 2 + 0.04, "builds", fontsize=8,
            ha="center", color="#555")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.set_title("System architecture", fontsize=12)
    plt.savefig(path, dpi=160, bbox_inches="tight")
    plt.close()


def mlflow_runs():
    mlflow.set_experiment("heart-disease")
    exp = mlflow.get_experiment_by_name("heart-disease")
    df = mlflow.search_runs(experiment_ids=[exp.experiment_id],
                            order_by=["start_time DESC"])
    if df.empty:
        sys.exit("No MLflow runs found. Run:  python -m src.train")
    return df, df.drop_duplicates(subset="tags.mlflow.runName")


def main():
    os.makedirs(os.path.join(ROOT, "reports"), exist_ok=True)
    if not os.path.isdir(EDA):
        sys.exit("EDA outputs missing. Run:  python -m src.eda")
    raw = load_raw()
    df = load_clean()
    all_runs, latest = mlflow_runs()
    meta_path = os.path.join(ROOT, "models", "metadata.json")
    meta = json.load(open(meta_path)) if os.path.exists(meta_path) else {}
    n_tests = sum(len(re.findall(r"^def test_", open(f).read(), re.M))
                  for f in glob.glob(os.path.join(ROOT, "tests", "test_*.py")))
    pos = int(df["target"].sum())
    neg = len(df) - pos
    arch = os.path.join(ROOT, "reports", "architecture.png")
    draw_architecture(arch)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    t = doc.add_heading("MLOps Assignment 01", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ("End-to-End ML Model Development, CI/CD and Production Deployment",
                 "Course: Machine Learning Operations (MLOps) AIMLCZG523, BITS Pilani",
                 f"Student: {STUDENT_NAME} ({STUDENT_ID})",
                 f"Code repository: {REPO_URL}"):
        c = doc.add_paragraph(line)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("1. Problem statement and approach", 1)
    p(doc, "The goal is to predict the presence of heart disease from patient health "
           "data and to serve the model as a containerised, monitored API deployed on "
           "Kubernetes. The work follows MLOps practice: a reproducible pipeline, "
           "experiment tracking, automated tests, CI, containerisation, orchestrated "
           "deployment and monitoring.")
    bullets(doc, [
        "Data: UCI Heart Disease (Cleveland) dataset, downloaded by a script.",
        "Models: Logistic Regression and Random Forest, tuned with grid search and "
        "5-fold stratified cross-validation, tracked in MLflow.",
        "Serving: FastAPI with /predict, /health and /metrics, packaged in Docker.",
        "Delivery: GitHub Actions (lint, tests, training, Docker build); "
        "Kubernetes Deployment and LoadBalancer Service on Minikube.",
        "Observability: request logging, Prometheus metrics and a Grafana dashboard.",
    ])

    doc.add_heading("2. Setup and installation", 1)
    p(doc, "Everything runs from a clean checkout. In the BITS Prayogshala lab "
           "(Docker available, Python 3.12):")
    code(doc, f"git clone {REPO_URL}.git\ncd mlops-heart-disease\n"
              "bash scripts/run_all.sh        # install, data, EDA, train, lint, tests, "
              "docker build/run, API smoke test\n"
              "bash scripts/monitoring_up.sh  # Prometheus + Grafana\n"
              "bash scripts/k8s_deploy.sh     # Minikube deployment")
    p(doc, "Each script runs pre-flight checks, stops with a clear ERROR message on "
           "failure and writes a log file to logs/. Manual equivalents:")
    code(doc, "pip install -r requirements.txt\npython -m src.download_data\n"
              "python -m src.eda\npython -m src.train\nmlflow ui\npython -m pytest -v\n"
              "docker build -t heart-api:latest .\n"
              "docker run -d -p 8000:8000 --name heart-api heart-api:latest")

    doc.add_heading("3. Data acquisition and exploratory data analysis", 1)
    p(doc, f"The dataset has {raw.shape[0]} patients and {len(FEATURES)} features plus "
           "a target. src/download_data.py fetches processed.cleveland.data from the "
           "UCI repository. Cleaning (src/data.py) binarises the original 0-4 target "
           "(0 = no disease, 1-4 = disease), removes duplicates and keeps the "
           "cleaned data in data/heart_clean.csv.")
    miss = raw.isna().sum()
    miss = miss[miss > 0]
    p(doc, "Missing values: " + (", ".join(f"{k} ({v})" for k, v in miss.items())
                                 if len(miss) else "none") +
      ". They are imputed inside the model pipeline (median for numeric features, "
      "most-frequent for categorical features), so imputation is learned only from "
      "training folds and applied identically at inference time.")
    p(doc, f"Class balance after cleaning: {pos} disease vs {neg} no disease "
           f"({100 * pos / len(df):.1f}% positive). The classes are reasonably "
           "balanced, so accuracy is meaningful, but precision, recall and ROC-AUC "
           "are also reported because the cost of a missed case is high.")
    for img, cap in [("histograms.png", "Figure 1: Distributions of numeric features"),
                     ("correlation_heatmap.png", "Figure 2: Correlation heatmap"),
                     ("class_balance.png", "Figure 3: Class balance"),
                     ("features_by_target.png", "Figure 4: Numeric features by target")]:
        path = os.path.join(EDA, img)
        if os.path.exists(path):
            picture(doc, path, cap, 5.2 if img == "class_balance.png" else 6.0)
    corr = df.corr()["target"].drop("target").sort_values(key=abs, ascending=False)
    p(doc, "Features most correlated with the target: " +
      ", ".join(f"{k} ({v:+.2f})" for k, v in corr.head(5).items()) + ".")

    doc.add_heading("4. Feature engineering and model development", 1)
    p(doc, "Preprocessing is a scikit-learn ColumnTransformer: numeric features "
           f"({', '.join(NUMERIC)}) are median-imputed and standardised; categorical "
           "features are most-frequent-imputed and one-hot encoded (unknown categories "
           "ignored at inference). The transformer and the classifier are fitted as "
           "one Pipeline, so the exact same preprocessing is applied in training and "
           "serving.")
    p(doc, "Training uses an 80/20 stratified train/test split (seed 42). Each model is "
           "tuned with GridSearchCV (5-fold stratified CV, scoring ROC-AUC) on the "
           "training split; the best configuration is then cross-validated for the "
           "full metric set and scored once on the held-out test split. The model "
           "to deploy is selected on cross-validation AUC, never on the test set.")
    table(doc, ["Model", "Search space"],
          [["Logistic Regression", "C in {0.01, 0.1, 1, 10}"],
           ["Random Forest", "n_estimators in {100, 300}; max_depth in {3, 5, None}"]])
    rows = []
    for _, r in latest.iterrows():
        rows.append([r["tags.mlflow.runName"].replace("_", " ").title()] + [
            f"{r.get('metrics.' + m, float('nan')):.3f}"
            for m in ("cv_accuracy", "cv_precision", "cv_recall",
                      "cv_roc_auc", "test_roc_auc")])
    table(doc, ["Model", "CV accuracy", "CV precision", "CV recall",
                "CV ROC-AUC", "Test ROC-AUC"], rows)
    if meta:
        cvs = {r["tags.mlflow.runName"]: r.get("metrics.cv_roc_auc")
               for _, r in latest.iterrows()}
        other = [k for k in cvs if k != meta["model"]]
        txt = (f"Selected model: {meta['model'].replace('_', ' ')} "
               f"(CV ROC-AUC {meta['cv_roc_auc']:.3f}, held-out test ROC-AUC "
               f"{meta['test_roc_auc']:.3f}).")
        if other:
            gap = meta["cv_roc_auc"] - cvs[other[0]]
            txt += (f" It leads {other[0].replace('_', ' ')} by {gap:.3f} CV "
                    "ROC-AUC")
            if gap < 0.02 and meta["model"] == "logistic_regression":
                txt += (", a negligible difference, so the simpler and more "
                        "interpretable linear model is preferred.")
            else:
                txt += "."
        p(doc, txt)
    shot(doc, "08_train_output.png", "Figure 5: Training output")

    doc.add_heading("5. Experiment tracking (MLflow)", 1)
    p(doc, "Every training run is logged to MLflow under the experiment "
           "'heart-disease': best hyper-parameters, cross-validated metrics (accuracy, "
           "precision, recall, ROC-AUC), test ROC-AUC, the ROC curve and confusion "
           "matrix plots, and the fitted model. Each execution of the training script "
           "adds one run per model, so repeated executions during development appear "
           "as several runs; the table above lists the latest run of each model.")
    prm = [c for c in latest.columns if c.startswith("params.")]
    rows = []
    for _, r in latest.iterrows():
        used = {c.replace("params.", ""): r[c] for c in prm
                if str(r[c]) not in ("None", "nan")}
        name = r["tags.mlflow.runName"].replace("_", " ").title()
        rows.append([name, ", ".join(f"{k}={v}" for k, v in used.items())])
    table(doc, ["Run", "Best parameters"], rows)
    shot(doc, "01_mlflow_runs.png",
         "Figure 6: MLflow comparison of the two models "
         "(left column: random forest, right column: logistic regression)")
    shot(doc, "02_mlflow_run_detail.png", "Figure 7: MLflow run with parameters, metrics and plots")

    doc.add_heading("6. Model packaging and reproducibility", 1)
    bullets(doc, [
        "The final model is the complete Pipeline (preprocessing + classifier), saved "
        "with joblib as models/model.joblib and logged to MLflow.",
        "requirements.txt lists all dependencies; every script installs from it.",
        "Random seeds are fixed (42) for the split, cross-validation and forest.",
        "The Docker image downloads the data and trains the model during the build, "
        "so training and serving always use identical library versions.",
    ])

    doc.add_heading("7. CI/CD pipeline and automated testing", 1)
    p(doc, f"The repository contains {n_tests} pytest tests covering data cleaning "
           "(target binarisation), the preprocessing pipeline (missing values and "
           "output shape), model training and probability output, and the API "
           "(health, metrics and prediction endpoints, including input validation).")
    p(doc, "The GitHub Actions workflow (.github/workflows/ci.yml) runs on every push "
           "and pull request: install dependencies, flake8 lint, dataset download, "
           "model training, pytest, Docker image build, and upload of the test "
           "report, trained model and cleaned data as artifacts. The pipeline fails "
           "on any lint, test or training error and the logs are kept per run.")
    p(doc, "The first workflow run failed during training: a newer MLflow release "
           "refused to save the scikit-learn pipeline in its default format. The "
           "failure was visible in the run logs, was reproduced locally, fixed by "
           "saving the model in cloudpickle format, and the following runs passed. "
           "This is the kind of defect the pipeline exists to catch.")
    shot(doc, "03_github_actions.png",
         "Figure 8: GitHub Actions run history "
         "(the latest runs pass; the first run failed and was fixed in run 2)")

    doc.add_heading("8. Containerisation", 1)
    p(doc, "The Dockerfile uses python:3.12-slim, installs requirements, copies the "
           "source, trains the model inside the image and starts uvicorn on port "
           "8000. The API accepts the 13 patient features as JSON at POST /predict and "
           "returns the predicted class and the confidence (probability of disease); "
           "invalid input returns HTTP 422. Sample request:")
    code(doc, 'curl -X POST localhost:8000/predict -H "Content-Type: application/json" '
              "-d @sample_input.json\n"
              '-> {"prediction": ..., "confidence": ...}')
    p(doc, "An early build used Python 3.10 while the model had been trained on "
           "Python 3.12 with a newer scikit-learn, which made /predict return HTTP "
           "500 inside the container. Training inside the image removed this "
           "train/serve mismatch.")
    shot(doc, "04_docker_build_run.png",
         "Figure 9: Running containers with /health and /predict responses")

    doc.add_heading("9. Production deployment (Kubernetes)", 1)
    p(doc, "The API is deployed on a local Minikube cluster (Docker driver) using "
           "k8s/deployment.yaml (2 replicas, readiness and liveness probes on /health, "
           "resource requests and limits) and k8s/service.yaml (Service of type "
           "LoadBalancer, port 80 to container port 8000). scripts/k8s_deploy.sh "
           "starts the cluster, loads the image, applies the manifests, waits for the "
           "rollout and verifies /health and /predict through the Service. On "
           "Minikube the LoadBalancer receives an external IP through 'minikube "
           "tunnel'.")
    p(doc, "Access instructions: the API is not hosted on a public URL. For local "
           "testing, run scripts/run_all.sh (container on http://localhost:8000, "
           "interactive docs at /docs) or scripts/k8s_deploy.sh (Kubernetes, then "
           "'minikube service heart-api --url'). The same model can also be scored "
           "from the command line with 'python -m src.predict sample_input.json'.")
    shot(doc, "05_k8s_deploy.png", "Figure 10: Kubernetes rollout and resources")

    doc.add_heading("10. Monitoring and logging", 1)
    p(doc, "The API logs every request (input and result) and every failure with a "
           "stack trace. It exposes Prometheus metrics at /metrics: "
           "predict_requests_total (labelled by status) and a latency histogram "
           "predict_latency_seconds. docker-compose in monitoring/ starts Prometheus "
           "(scraping the API every 10 s) and Grafana with a provisioned datasource "
           "and a 'Heart Disease API' dashboard showing total predictions, errors, "
           "request rate and p95 latency. scripts/monitoring_up.sh generates traffic "
           "and verifies that Prometheus received the data.")
    shot(doc, "06_prometheus_targets.png", "Figure 11: Prometheus target UP")
    shot(doc, "07_grafana_dashboard.png", "Figure 12: Grafana dashboard")

    doc.add_heading("11. Architecture", 1)
    picture(doc, arch, "Figure 13: System architecture", 6.3)

    doc.add_heading("12. Conclusion and links", 1)
    p(doc, "The project delivers a reproducible pipeline from raw data to a monitored "
           "model service on Kubernetes, with automated quality gates. Possible "
           "extensions: model registry and promotion, data-drift monitoring, "
           "canary releases and automated retraining.")
    p(doc, f"Code repository: {REPO_URL}")
    doc.save(OUT)
    print(f"Report written to {os.path.abspath(OUT)}")
    if missing:
        print("Screenshots not found in screenshots/ (placeholders inserted):")
        for m in missing:
            print("  -", m)


if __name__ == "__main__":
    main()
